"""End-to-end API tests: auth, upload, risk, scenarios, chat, alerts, access control."""


def test_health(client):
    assert client.get("/api/health").json()["status"] == "ok"


def test_signup_login_me(client, make_user):
    client.post("/api/auth/signup", json={"email": "Login@Example.com", "password": "password123"})
    r = client.post("/api/auth/login", json={"email": "login@example.com", "password": "password123"})
    assert r.status_code == 200
    headers = {"Authorization": f"Bearer {r.json()['access_token']}"}
    assert client.get("/api/auth/me", headers=headers).json()["email"] == "login@example.com"


def test_wrong_password_and_duplicate_email(client):
    client.post("/api/auth/signup", json={"email": "dup@example.com", "password": "password123"})
    assert client.post("/api/auth/signup", json={"email": "dup@example.com", "password": "password123"}).status_code == 409
    assert client.post("/api/auth/login", json={"email": "dup@example.com", "password": "nope-nope"}).status_code == 401


def test_requires_login(client):
    assert client.get("/api/portfolio/risk").status_code == 401
    assert client.get("/api/portfolio/risk", headers={"Authorization": "Bearer junk"}).status_code == 401


def test_short_password_rejected(client):
    assert client.post("/api/auth/signup", json={"email": "s@example.com", "password": "short"}).status_code == 422


def test_upload_and_risk_report(client, user_with_portfolio):
    headers, _ = user_with_portfolio
    holdings = client.get("/api/portfolio/holdings", headers=headers).json()
    assert len(holdings) == 10
    assert {h["sector"] for h in holdings} >= {"Information Technology", "Financial Services"}

    r = client.get("/api/portfolio/risk", headers=headers)
    assert r.status_code == 200
    rep = r.json()
    m = rep["metrics"]
    assert 0 < m["annual_volatility"] < 1
    assert m["max_drawdown"] <= 0
    assert m["var_95_amount"] > 0
    assert m["cvar_95_amount"] >= m["var_95_amount"]
    assert abs(sum(h["weight"] for h in rep["holdings"]) - 1) < 1e-9
    assert abs(sum(h["risk_share"] for h in rep["holdings"]) - 1) < 1e-6
    assert len(rep["history"]) > 200
    assert len(rep["correlation_matrix"]["symbols"]) == 10


def test_upload_bad_file(client, make_user):
    headers, _ = make_user()
    r = client.post("/api/portfolio/upload", headers=headers, files={"file": ("x.csv", b"a,b\n1,2", "text/csv")})
    assert r.status_code == 422


def test_add_edit_delete_holding(client, make_user):
    headers, _ = make_user()
    r = client.post("/api/portfolio/holdings", headers=headers, json={"symbol": "infy", "quantity": 5, "avg_price": 1500})
    assert r.status_code == 201
    h = r.json()
    assert h["symbol"] == "INFY" and h["sector"] == "Information Technology"
    assert client.post("/api/portfolio/holdings", headers=headers,
                       json={"symbol": "INFY", "quantity": 1, "avg_price": 1}).status_code == 409
    r = client.put(f"/api/portfolio/holdings/{h['id']}", headers=headers,
                   json={"symbol": "INFY", "quantity": 7, "avg_price": 1400})
    assert r.json()["quantity"] == 7
    assert client.delete(f"/api/portfolio/holdings/{h['id']}", headers=headers).status_code == 204


def test_users_cannot_touch_each_others_data(client, user_with_portfolio, make_user):
    owner_headers, _ = user_with_portfolio
    other_headers, _ = make_user()
    holding_id = client.get("/api/portfolio/holdings", headers=owner_headers).json()[0]["id"]
    assert client.delete(f"/api/portfolio/holdings/{holding_id}", headers=other_headers).status_code == 404
    assert client.get("/api/portfolio/holdings", headers=other_headers).json() == []
    assert client.get("/api/portfolio/risk", headers=other_headers).status_code == 404


def test_scenarios(client, user_with_portfolio):
    headers, _ = user_with_portfolio
    r = client.post("/api/scenarios/sector", headers=headers, json={"sector": "IT", "drop_pct": 15})
    assert r.status_code == 200 and r.json()["total_impact"] < 0
    r = client.post("/api/scenarios/market", headers=headers, json={"drop_pct": 10})
    assert r.status_code == 200 and r.json()["total_impact"] < 0
    events = client.get("/api/scenarios/events").json()
    for e in events:
        r = client.post("/api/scenarios/historical", headers=headers, json={"event": e["key"]})
        assert r.status_code == 200, r.text
    r = client.post("/api/scenarios/sector", headers=headers, json={"sector": "Steel", "drop_pct": 15})
    assert r.status_code == 422


def test_chat_offline(client, user_with_portfolio):
    headers, _ = user_with_portfolio
    r = client.post("/api/chat", headers=headers, json={"message": "What if IT falls 15%?"})
    assert r.status_code == 200
    body = r.json()
    assert body["mode"] == "offline"
    assert any(c["tool"] == "run_sector_scenario" for c in body["tool_calls"])
    assert "₹" in body["answer"]


def test_alerts(client, user_with_portfolio):
    headers, _ = user_with_portfolio
    first = client.post("/api/alerts/check", headers=headers).json()["new_alerts"]
    second = client.post("/api/alerts/check", headers=headers).json()["new_alerts"]
    assert second == 0  # same alerts are never created twice
    alerts = client.get("/api/alerts", headers=headers).json()
    assert len(alerts) == first


def test_settings_update(client, make_user):
    headers, _ = make_user()
    r = client.patch("/api/auth/me", headers=headers, json={"telegram_chat_id": "12345", "name": "Ankit"})
    assert r.json()["telegram_chat_id"] == "12345"
    assert client.post("/api/alerts/test-telegram", headers=headers).status_code == 400  # no bot token in tests


def test_admin_rbac(client, make_user, user_with_portfolio):
    user_headers, user = make_user()
    assert user["role"] == "user"
    assert client.get("/api/admin/stats", headers=user_headers).status_code == 403

    admin_headers, admin = make_user("admin@example.com")
    assert admin["role"] == "admin"
    stats = client.get("/api/admin/stats", headers=admin_headers).json()
    assert stats["users"] >= 2 and stats["users_with_portfolio"] >= 1


def test_cron_endpoint(client, user_with_portfolio, monkeypatch):
    from app.routers import cron

    monkeypatch.setattr(cron.settings, "cron_secret", "")
    assert client.post("/api/cron/daily-alerts").status_code == 404  # disabled when no secret is set

    monkeypatch.setattr(cron.settings, "cron_secret", "s3cret-value")
    assert client.post("/api/cron/daily-alerts", headers={"X-Cron-Secret": "wrong"}).status_code == 403
    r = client.post("/api/cron/daily-alerts", headers={"X-Cron-Secret": "s3cret-value"})
    assert r.status_code == 200 and r.json()["users_checked"] >= 1


def test_health_reports_database(client):
    body = client.get("/api/health").json()
    assert body["database"] == "ok"


def test_excel_report(client, user_with_portfolio):
    import io

    from openpyxl import load_workbook

    headers, _ = user_with_portfolio
    r = client.get("/api/portfolio/report.xlsx", headers=headers)
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Holdings", "Sectors", "Scenarios", "Correlation", "Daily history"]
    holdings = wb["Holdings"]
    assert holdings["A1"].value == "Symbol" and holdings.max_row == 11  # header + 10 stocks
    measures = [c.value for c in wb["Summary"]["A"]]
    assert "1-day Value at Risk (95%)" in measures and "VaR backtest" in measures
    scen = [c.value for c in wb["Scenarios"]["A"]][1:]
    assert "NIFTY 50 falls 10%" in scen and "COVID crash (2020)" in scen


def test_report_includes_outlook(client, user_with_portfolio):
    headers, _ = user_with_portfolio
    o = client.get("/api/portfolio/risk", headers=headers).json()["outlook"]
    assert o["month"]["paths"] == 10000 and o["month"]["bad_month"] < 0
    assert o["volatility"]["level"] in {"low", "normal", "elevated", "high"}
    assert o["var_backtest"]["days_tested"] == 250
